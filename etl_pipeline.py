import os
import io
import time
import sqlite3
import pandas as pd
from datetime import datetime, timedelta
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaFileUpload
import openpyxl

# ==========================================
# ⚙️ 구글 드라이브 & 시트 실제 ID 설정 (검증 완료)
# ==========================================

TOP_FOLDER_ID = '1UlsDUOZv3QPp19M_vMNptLiDZjEHPPUw'       # 대시보드 업로드
B2C_FOLDER_ID = '1ArGfyeVpZDJYUrdlGNSCrhj734JrqGW9'        # B2C 폴더
INBOUND_FOLDER_ID = '1BzKHxqaUrTFDubvJ7wnfXZqzNHaEvJjp'    # 입고 폴더
PROCESSED_FOLDER_ID = '1RiUOVDt8VEgOnePr_bje-ZPuzqlTYOXZ'  # 처리완료 폴더
DB_FOLDER_ID = '1jfi8ls7PWm9BWQUZwVYj9km5zAEuUKg9'         # DB 전용 폴더

DEFAULT_IB_SHEET_ID = '1j3yHXjpOpdYRBI_dFP6TBBG3Q3_vgbMAi3DW4po0SD0'
DB_INBOUND_PATH = 'wms_inbound.db'

# PPH/MH 구글 시트 실제 ID
SHEET_ID_375 = "1x8zWC0BZRs4ji9N8-CLPb9RutbegxcVh0BXNAlhGS6g"
SHEET_ID_XFC = "1IGgpr8LezBs4hCySj_W34HSyHFhN6lt13uYSt3gbmqbs"


def get_drive_service(creds_dict):
    creds = Credentials.from_service_account_info(
        creds_dict,
        scopes=[
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/spreadsheets.readonly'
        ]
    )
    return build('drive', 'v3', credentials=creds)


def get_sheets_service(creds_dict):
    creds = Credentials.from_service_account_info(
        creds_dict,
        scopes=['https://www.googleapis.com/auth/spreadsheets.readonly']
    )
    return build('sheets', 'v4', credentials=creds)


def safe_drive_list(service, query):
    try:
        res = service.files().list(
            q=query, fields="files(id, name, mimeType, parents)",
            supportsAllDrives=True, includeItemsFromAllDrives=True
        ).execute()
        return res.get('files', [])
    except Exception:
        return []


def get_or_create_subfolder(service, parent_id, folder_name):
    """'처리완료/YYYY-MM' 및 '처리완료/중복파일' 하위 폴더 자동 생성/조회"""
    q = f"'{parent_id}' in parents and name = '{folder_name}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
    items = safe_drive_list(service, q)
    if items:
        return items[0]['id']

    meta = {
        'name': folder_name,
        'mimeType': 'application/vnd.google-apps.folder',
        'parents': [parent_id]
    }
    try:
        folder = service.files().create(body=meta, fields='id', supportsAllDrives=True).execute()
        return folder.get('id')
    except Exception:
        return parent_id


def download_db_from_drive(service, db_filename):
    try:
        query = f"'{DB_FOLDER_ID}' in parents and name = '{db_filename}' and trashed = false"
        files = safe_drive_list(service, query)

        if files:
            file_id = files[0]['id']
            request = service.files().get_media(fileId=file_id)
            with open(db_filename, 'wb') as f:
                downloader = MediaIoBaseDownload(f, request)
                done = False
                while not done:
                    _, done = downloader.next_chunk()
            return True
    except Exception as e:
        print(f"DB Download Warning ({db_filename}): {e}")
    return False


def upload_db_to_drive(service, db_filename):
    if not os.path.exists(db_filename):
        return
    try:
        query = f"'{DB_FOLDER_ID}' in parents and name = '{db_filename}' and trashed = false"
        files = safe_drive_list(service, query)

        media = MediaFileUpload(db_filename, mimetype='application/x-sqlite3', resumable=True)

        if files:
            file_id = files[0]['id']
            service.files().update(fileId=file_id, media_body=media, supportsAllDrives=True).execute()
        else:
            file_metadata = {'name': db_filename, 'parents': [DB_FOLDER_ID]}
            service.files().create(body=file_metadata, media_body=media, fields='id', supportsAllDrives=True).execute()
    except Exception as e:
        print(f"DB Upload Error ({db_filename}): {e}")


def list_files_in_folder(service, folder_id):
    try:
        query = f"'{folder_id}' in parents and trashed = false"
        files = safe_drive_list(service, query)
        return [f for f in files if not f['name'].lower().endswith('.db')]
    except Exception as e:
        print(f"Folder list error ({folder_id}): {e}")
        return []


def read_excel_fast(fh):
    try:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        target_sheet = wb.active

        for sheet_name in wb.sheetnames:
            s = wb[sheet_name]
            for row in list(s.iter_rows(max_row=5, values_only=True)):
                row_str = " ".join([str(v) for v in row if v is not None])
                if any(k in row_str for k in ['송장 번호', '송장번호', '마감 일시', '마감일시', 'SKU명', '입고번호', '입고 번호']):
                    target_sheet = s
                    break
            else:
                continue
            break

        rows = target_sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        
        clean_headers = []
        counts = {}
        for h in headers:
            h_str = str(h).strip() if h is not None else "Unnamed"
            counts[h_str] = counts.get(h_str, 0) + 1
            clean_headers.append(f"{h_str}_{counts[h_str]}" if counts[h_str] > 1 else h_str)

        data = [r for r in rows if any(v is not None for v in r)]
        df = pd.DataFrame(data, columns=clean_headers)
        wb.close()
        return df
    except Exception:
        fh.seek(0)
        return pd.read_excel(fh, engine='openpyxl')


def update_inbound_plt_box_from_sheets(service, sheets_service, conn_ib):
    if not service or not sheets_service:
        return 0, "서비스 계정 권한 없음"

    matched_count = 0
    err_msg = ""

    try:
        sheet_metadata = sheets_service.spreadsheets().get(spreadsheetId=DEFAULT_IB_SHEET_ID).execute()
        sheets = sheet_metadata.get('sheets', [])
        
        if sheets:
            ib_sheet_title = None
            for s in sheets:
                title = s.get('properties', {}).get('title', '').strip()
                title_clean = title.replace(" ", "")
                if title_clean == '입고' or '입고' in title_clean or title_clean.upper() == 'IB':
                    ib_sheet_title = title
                    break

            if not ib_sheet_title:
                ib_sheet_title = sheets[0].get('properties', {}).get('title', 'Sheet1')

            range_name = f"'{ib_sheet_title}'!A:Z"
            result = sheets_service.spreadsheets().values().get(spreadsheetId=DEFAULT_IB_SHEET_ID, range=range_name).execute()
            values = result.get('values', [])

            if values and len(values) >= 2:
                header_row_idx = 0
                for r_i, r_data in enumerate(values[:10]):
                    r_str = " ".join([str(v) for v in r_data])
                    if '작업번호' in r_str or '입고번호' in r_str or 'PLT' in r_str:
                        header_row_idx = r_i
                        break

                headers = [str(h).strip() for h in values[header_row_idx]]

                job_no_idx, plt_idx, box_idx, pajok_idx = -1, -1, -1, -1
                for idx, h in enumerate(headers):
                    h_clean = h.replace(" ", "").upper()
                    if '작업번호' in h_clean or '입고번호' in h_clean: job_no_idx = idx
                    elif h_clean == 'PLT': plt_idx = idx
                    elif h_clean == 'BOX': box_idx = idx
                    elif '파적' in h_clean and 'BOX' in h_clean: pajok_idx = idx

                if job_no_idx != -1:
                    update_tuples = []
                    for row in values[header_row_idx + 1:]:
                        if len(row) > job_no_idx:
                            job_no = str(row[job_no_idx]).strip()
                            if not job_no or job_no.lower() == 'none' or job_no == '작업번호':
                                continue

                            def parse_val(row_data, idx_pos):
                                if idx_pos != -1 and len(row_data) > idx_pos:
                                    v_str = str(row_data[idx_pos]).replace(',', '').strip()
                                    try: return float(v_str)
                                    except Exception: return 0.0
                                return 0.0

                            plt_val = parse_val(row, plt_idx)
                            box_val = parse_val(row, box_idx)
                            pajok_val = parse_val(row, pajok_idx)

                            if plt_val > 0 or box_val > 0 or pajok_val > 0:
                                update_tuples.append((plt_val, box_val, pajok_val, job_no))

                    if update_tuples:
                        conn_ib.executemany("""
                        UPDATE inbound_summary
                        SET PLT수 = ?, BOX수 = ?, 파적BOX수 = ?
                        WHERE 입고번호 = ?
                        """, update_tuples)
                        conn_ib.commit()
                        matched_count = len(update_tuples)
    except Exception as e_sheet:
        err_msg = str(e_sheet)

    return matched_count, err_msg


# ==========================================
# 📊 구글 시트 M/H 데이터 파싱 모듈 (요약행 원천 차단 + 실작업 시간 계산)
# ==========================================

def parse_mh_375(sheets_service):
    """
    375 센터: '2026PPH' 정제 시트 직결 수집
    A열: 날자, B열: 시작, C열: 종료, D열: 작업시간, E열: M/H, F열: PPH
    """
    rows_to_save = []
    try:
        # 새로 만드신 2026PPH 시트의 A2:F 전체 범위를 직접 수집
        range_name = "'2026PPH'!A2:F"
        result = sheets_service.spreadsheets().values().get(
            spreadsheetId=SHEET_ID_375, range=range_name
        ).execute()
        values = result.get('values', [])

        curr_year = datetime.now().year

        for row in values:
            if not row or len(row) < 5:
                continue

            raw_date = str(row[0]).strip()
            start_time = str(row[1]).strip()
            mh_raw = str(row[4]).strip()

            if not raw_date or not start_time or ":" not in start_time or not mh_raw:
                continue

            # 날짜 형식 자동 정밀 변환 (2026. 10. 15 / 10/15 목 / 2026-10-15 모두 대응)
            date_str = None
            if "." in raw_date:
                parts = [p.strip() for p in raw_date.split(".") if p.strip()]
                if len(parts) >= 3 and parts[0].isdigit() and parts[1].isdigit():
                    y = int(parts[0])
                    m = int(parts[1])
                    d = int(parts[2].split()[0])
                    date_str = f"{y:04d}-{m:02d}-{d:02d}"
            elif "-" in raw_date:
                parts = [p.strip() for p in raw_date.split("-") if p.strip()]
                if len(parts) >= 3 and parts[0].isdigit() and parts[1].isdigit():
                    y = int(parts[0])
                    m = int(parts[1])
                    d = int(parts[2].split()[0])
                    date_str = f"{y:04d}-{m:02d}-{d:02d}"
            elif "/" in raw_date:
                clean_date_part = raw_date.split()[0].strip()
                md = clean_date_part.split("/")
                if len(md) == 2 and md[0].isdigit() and md[1].isdigit():
                    date_str = f"{curr_year}-{int(md[0]):02d}-{int(md[1]):02d}"

            if not date_str:
                continue

            try:
                # 서식 공백 대비 split()[0] 처리
                clean_mh = float(str(mh_raw).split()[0].replace(",", "").replace("-", "0"))
                if clean_mh <= 0:
                    continue

                h_int = int(start_time.split(":")[0])
                time_label = f"{h_int:02d}:00 (D)" if h_int >= 6 else f"{h_int:02d}:00 (D+1)"

                rows_to_save.append({
                    '영업마감일자': date_str,
                    '시간대': time_label,
                    '센터그룹': '375',
                    'MH': clean_mh
                })
            except Exception:
                continue
    except Exception as e:
        print(f"2026PPH 수집 에러: {e}")
    return rows_to_save
    

def parse_mh_xfc(sheets_service):
  """XFC 센터: (XFC센터)출고Dashboard -> 'PPH-(일)(자동화)'

  A열: 날짜(0), G열: 시작(6), H열: 종료(7), I열: 작업시간(8), J열: M/H(9)
  """
  rows_to_save = []
  try:
    # 1. 탭 이름 자동 탐색 ('PPH-(일)(자동화)' 포함 여부 확인)
    meta = (
        sheets_service.spreadsheets()
        .get(spreadsheetId=SHEET_ID_XFC)
        .execute()
    )
    target_tab = "PPH-(일)(자동화)"
    for s in meta.get("sheets", []):
      t_title = s.get("properties", {}).get("title", "")
      if "PPH-(일)(자동화)" in t_title.replace(" ", ""):
        target_tab = t_title
        break

    # 2. A열(날짜)부터 J열(M/H)까지만 수집 (인덱스 9번까지만 필요하므로 21칸까지 갈 필요 없음)
    range_name = f"'{target_tab}'!A6:J10000"
    result = (
        sheets_service.spreadsheets()
        .values()
        .get(spreadsheetId=SHEET_ID_XFC, range=range_name)
        .execute()
    )
    values = result.get("values", [])

    for row in values:
      # 최소 J열(인덱스 9)까지 존재하는 행만 확인
      if not row or len(row) < 10:
        continue

      raw_date = str(row[0]).strip()
      start_time = str(row[6]).strip()
      work_time_raw = str(row[8]).strip()
      mh_raw = str(row[9]).strip()  # ⭐ J열에 이미 완성된 M/H 값 직접 수집

      if not raw_date or ":" not in start_time or not mh_raw:
        continue

      date_clean = raw_date[:10]  # '2026-10-15(목)' -> '2026-10-15'

      try:
        # 작업시간이 0.1h ~ 1.5h 사이인 실작업 시간대만 인정
        work_time = float(work_time_raw.replace(",", "").replace("-", "0"))
        if work_time <= 0.1 or work_time > 1.5:
          continue

        clean_mh = float(mh_raw.replace(",", "").replace("-", "0"))
        if clean_mh <= 0:
          continue

        h_int = int(start_time.split(":")[0])
        time_label = (
            f"{h_int:02d}:00 (D)" if h_int >= 6 else f"{h_int:02d}:00 (D+1)"
        )

        rows_to_save.append({
            "영업마감일자": date_clean,
            "시간대": time_label,
            "센터그룹": "XFC",
            "MH": clean_mh,
        })
      except Exception:
        continue
  except Exception as e:
    print(f"XFC MH 수집 에러: {e}")
  return rows_to_save
    

def sync_mh_to_db(sheets_service, drive_service=None, db_path="wms_b2c.db"):
    list_375 = parse_mh_375(sheets_service)
    list_xfc = parse_mh_xfc(sheets_service)
    all_mh = list_375 + list_xfc

    if all_mh:
        df_mh = pd.DataFrame(all_mh)
        # 시간대별 고유 집계 (30분 단위 등이 1시간 내 2건 있을 경우 정상 합산, 요약행은 완전 배제됨)
        df_mh_agg = df_mh.groupby(['영업마감일자', '시간대', '센터그룹'])['MH'].sum().reset_index()
        target_dbs = [f for f in os.listdir('.') if f.startswith('wms_b2c') and f.endswith('.db')]
        if not target_dbs: target_dbs = [db_path]
            
        for d_path in set(target_dbs):
            try:
                conn = sqlite3.connect(d_path, timeout=10)
                conn.execute("""
                CREATE TABLE IF NOT EXISTS center_mh_hourly (
                    영업마감일자 TEXT,
                    시간대 TEXT,
                    센터그룹 TEXT,
                    MH REAL,
                    PRIMARY KEY (영업마감일자, 시간대, 센터그룹)
                )
                """)
                conn.execute("DELETE FROM center_mh_hourly")
                df_mh_agg.to_sql("center_mh_hourly", conn, if_exists="append", index=False)
                conn.commit()
                conn.close()
                if drive_service:
                    upload_db_to_drive(drive_service, d_path)
            except Exception as e:
                print(f"MH 테이블 DB 저장 에러 ({d_path}): {e}")
        return len(df_mh_agg)
    return 0


# ==========================================
# 🚀 메인 동기화 프로세스 실행 함수
# ==========================================

def process_and_update(service, sheets_service=None, progress_callback=None):
    if not os.path.exists(DB_INBOUND_PATH):
        download_db_from_drive(service, DB_INBOUND_PATH)

    conn_ib = sqlite3.connect(DB_INBOUND_PATH, timeout=30)
    conn_ib.execute("""
    CREATE TABLE IF NOT EXISTS inbound_summary (
        영업마감일자 TEXT, 센터 TEXT, 고객사 TEXT, 상태 TEXT, 입고번호 TEXT,
        입고방법 TEXT, SKU명 TEXT, 바코드 TEXT, 소비기한 TEXT, 로트 TEXT,
        기본로케이션 TEXT, 예정수량 INTEGER, 요청SKU수량 INTEGER, 총예정수량 INTEGER,
        총검수완료수량 INTEGER, PLT수 REAL, BOX수 REAL, 파적BOX수 REAL, 등록일시 TEXT, 변경자 TEXT,
        최종변경일시 TEXT, 입고완료일시 TEXT,
        PRIMARY KEY (영업마감일자, 센터, 고객사, 상태, 입고번호, SKU명, 바코드)
    )
    """)

    folder_mapping = [
        (INBOUND_FOLDER_ID, 'INBOUND'),
        (B2C_FOLDER_ID, 'B2C'),
        (TOP_FOLDER_ID, 'AUTO')
    ]

    all_target_files = []
    for f_id, category in folder_mapping:
        files = list_files_in_folder(service, f_id)
        for f in files:
            if f['name'].lower().endswith('.xlsx') or f['name'].lower().endswith('.csv'):
                f['category'] = category
                f['source_folder_id'] = f_id
                all_target_files.append(f)

    total_count = len(all_target_files)
    b2c_updated_months = set()
    inbound_updated = False

    duplicate_folder_id = get_or_create_subfolder(service, PROCESSED_FOLDER_ID, "중복파일")

    for idx, f in enumerate(all_target_files, 1):
        file_id, file_name = f['id'], f['name']
        category = f['category']
        src_folder = f['source_folder_id']

        if progress_callback:
            progress_callback(current=idx, total=total_count, filename=file_name, eta=0)

        try:
            request = service.files().get_media(fileId=file_id)
            fh = io.BytesIO()
            downloader = MediaIoBaseDownload(fh, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            fh.seek(0)

            if file_name.lower().endswith('.csv'):
                df = pd.read_csv(fh)
            else:
                df = read_excel_fast(fh)

            df_cols_no_space = [str(c).replace(" ", "").strip() for c in df.columns]
            is_inbound = (category == 'INBOUND') or any(k in "".join(df_cols_no_space) for k in ['입고번호', '총검수완료수량', '입고방법']) or ('입고요청서' in file_name)

            if is_inbound:
                col_map_inbound = {}
                for orig_c in df.columns:
                    clean_c = str(orig_c).replace(" ", "").strip()
                    if '상태' == clean_c: col_map_inbound[orig_c] = '상태'
                    elif '센터' == clean_c: col_map_inbound[orig_c] = '센터'
                    elif '고객사' == clean_c: col_map_inbound[orig_c] = '고객사'
                    elif '입고번호' in clean_c: col_map_inbound[orig_c] = '입고번호'
                    elif '입고방법' in clean_c: col_map_inbound[orig_c] = '입고방법'
                    elif 'SKU명' in clean_c or '상품명' in clean_c: col_map_inbound[orig_c] = 'SKU명'
                    elif '바코드' == clean_c: col_map_inbound[orig_c] = '바코드'
                    elif '소비기한' in clean_c or '유통기한' in clean_c: col_map_inbound[orig_c] = '소비기한'
                    elif '로트' in clean_c or 'LOT' in clean_c.upper(): col_map_inbound[orig_c] = '로트'
                    elif '로케이션' in clean_c: col_map_inbound[orig_c] = '기본로케이션'
                    elif '예정수량' == clean_c: col_map_inbound[orig_c] = '예정수량'
                    elif '요청SKU' in clean_c or '요청sku' in clean_c: col_map_inbound[orig_c] = '요청SKU수량'
                    elif '총예정수량' in clean_c: col_map_inbound[orig_c] = '총예정수량'
                    elif '총검수완료수량' in clean_c or '검수완료' in clean_c: col_map_inbound[orig_c] = '총검수완료수량'
                    elif '등록일시' in clean_c: col_map_inbound[orig_c] = '등록일시'
                    elif '변경자' in clean_c: col_map_inbound[orig_c] = '변경자'
                    elif '최종변경일시' in clean_c or '최종변경' in clean_c: col_map_inbound[orig_c] = '최종변경일시'
                    elif '입고완료일시' in clean_c or '입고완료' in clean_c: col_map_inbound[orig_c] = '입고완료일시'

                df_in = df.rename(columns=col_map_inbound)
                df_in = df_in.loc[:, ~df_in.columns.duplicated()]

                if '상태' not in df_in.columns: df_in['상태'] = '입고 완료'

                def parse_inbound_date(row):
                    st_val = str(row.get('상태', '')).replace(" ", "").strip()
                    raw_date = row.get('입고완료일시', None) if '입고완료' in st_val else row.get('최종변경일시', None)
                    if pd.isna(raw_date) or str(raw_date).strip() == '' or str(raw_date) == 'nan':
                        raw_date = row.get('등록일시', None)
                    try: return pd.to_datetime(raw_date).strftime('%Y-%m-%d')
                    except Exception: return datetime.now().strftime('%Y-%m-%d')

                df_in['영업마감일자'] = df_in.apply(parse_inbound_date, axis=1)

                for num_c in ['예정수량', '요청SKU수량', '총예정수량', '총검수완료수량']:
                    if num_c in df_in.columns:
                        df_in[num_c] = pd.to_numeric(df_in[num_c], errors='coerce').fillna(0)

                for _, row_in in df_in.iterrows():
                    conn_ib.execute("""
                    INSERT OR REPLACE INTO inbound_summary
                    (영업마감일자, 센터, 고객사, 상태, 입고번호, 입고방법, SKU명, 바코드, 소비기한, 로트,
                     기본로케이션, 예정수량, 요청SKU수량, 총예정수량, 총검수완료수량, PLT수, BOX수, 파적BOX수, 등록일시, 변경자,
                     최종변경일시, 입고완료일시)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0.0, 0.0, 0.0, ?, ?, ?, ?)
                    """, (
                        row_in.get('영업마감일자', ''), row_in.get('센터', ''), row_in.get('고객사', ''), row_in.get('상태', ''),
                        str(row_in.get('입고번호', '')), str(row_in.get('입고방법', '')), str(row_in.get('SKU명', '')), str(row_in.get('바코드', '')),
                        str(row_in.get('소비기한', '')), str(row_in.get('로트', '')), str(row_in.get('기본로케이션', '')),
                        int(row_in.get('예정수량', 0)), int(row_in.get('요청SKU수량', 0)), int(row_in.get('총예정수량', 0)), int(row_in.get('총검수완료수량', 0)),
                        str(row_in.get('등록일시', '')), str(row_in.get('변경자', '')), str(row_in.get('최종변경일시', '')), str(row_in.get('입고완료일시', ''))
                    ))
                conn_ib.commit()
                inbound_updated = True

                latest_date_str = df_in['영업마감일자'].max()
                target_month_str = latest_date_str[:7] if pd.notnull(latest_date_str) and len(str(latest_date_str)) >= 7 else datetime.now().strftime('%Y-%m')
                target_folder_id = get_or_create_subfolder(service, PROCESSED_FOLDER_ID, target_month_str)

                service.files().update(
                    fileId=file_id, addParents=target_folder_id, removeParents=src_folder,
                    supportsAllDrives=True, fields='id, parents'
                ).execute()

            else:
                col_map_b2c = {}
                for orig_c in df.columns:
                    clean_c = str(orig_c).replace(" ", "").strip()
                    if '상세' in clean_c or '이형' in clean_c or '해외' in clean_c: continue
                    
                    if clean_c == '송장번호' or clean_c == '송장' or clean_c == '운송장번호': col_map_b2c[orig_c] = '송장번호'
                    elif clean_c == '출고수량' or clean_c == '수량' or clean_c == '출고수': col_map_b2c[orig_c] = '총출고수량'
                    elif 'SKU' in clean_c or '상품명' in clean_c: col_map_b2c[orig_c] = 'SKU명'
                    elif clean_c == '바코드': col_map_b2c[orig_c] = '바코드'
                    elif '센터' in clean_c: col_map_b2c[orig_c] = '센터'
                    elif '고객사' in clean_c: col_map_b2c[orig_c] = '고객사'
                    elif '배송속성' in clean_c or '배송유형' in clean_c: col_map_b2c[orig_c] = '배송속성'
                    elif '판매플랫폼' in clean_c or '판매처' in clean_c: col_map_b2c[orig_c] = '판매처'
                    elif clean_c == '출고박스' or clean_c == '박스' or (('출고박스' in clean_c or '박스' in clean_c) and '바코드' not in clean_c and '사용' not in clean_c): col_map_b2c[orig_c] = '출고박스종류'
                    elif '마감일시' in clean_c or '마감일' in clean_c: col_map_b2c[orig_c] = '마감일시'
                    elif '마감자' == clean_c: col_map_b2c[orig_c] = '마감자'
                    elif '주문일시' in clean_c or '주문일' in clean_c: col_map_b2c[orig_c] = '주문일시'
                    elif '결제일시' in clean_c or '결제일' in clean_c: col_map_b2c[orig_c] = '결제일시'
                    elif '등록일시' in clean_c or '등록일' in clean_c: col_map_b2c[orig_c] = '등록일시'
                    elif '할당일시' in clean_c or '할당일' in clean_c: col_map_b2c[orig_c] = '할당일시'
                    elif '출력일시' in clean_c or '출력일' in clean_c: col_map_b2c[orig_c] = '출력일시'
                    elif '브랜드' in clean_c: col_map_b2c[orig_c] = '브랜드'
                    elif '배송계약' in clean_c or '태그' in clean_c: col_map_b2c[orig_c] = '배송계약태그'
                    elif '개별주문' in clean_c or '개별주문번호' in clean_c: col_map_b2c[orig_c] = '개별주문번호'
                    elif clean_c == '주문번호': col_map_b2c[orig_c] = '주문번호'
                    elif '피킹지시서' in clean_c and '총량' not in clean_c: col_map_b2c[orig_c] = '피킹지시서번호'
                    elif '품고추적' in clean_c or '품고' in clean_c: col_map_b2c[orig_c] = '품고추적번호'
                    elif clean_c == 'CS' or clean_c == 'cs': col_map_b2c[orig_c] = 'CS'

                df_b2c_f = df.rename(columns=col_map_b2c)
                df_b2c_f = df_b2c_f.loc[:, ~df_b2c_f.columns.duplicated()].ffill()

                date_col_name = None
                for c in ['마감일시', '주문일시', '등록일시']:
                    if c in df_b2c_f.columns:
                        date_col_name = c
                        break

                if date_col_name is not None:
                    raw_date_data = df_b2c_f[date_col_name]
                    df_b2c_f['dt_temp'] = pd.to_datetime(raw_date_data, errors='coerce')
                    df_b2c_f['영업마감일자'] = (df_b2c_f['dt_temp'] - pd.Timedelta(hours=6)).dt.strftime('%Y-%m-%d')
                else:
                    df_b2c_f['영업마감일자'] = datetime.now().strftime('%Y-%m-%d')

                df_b2c_f['영업마감일자'] = df_b2c_f['영업마감일자'].fillna(datetime.now().strftime('%Y-%m-%d'))

                req_cols_b2c = [
                    '센터', '고객사', '배송속성', '판매처', '출고박스종류', 'SKU명', '바코드', '송장번호',
                    '마감일시', '마감자', '주문일시', '결제일시', '등록일시', '할당일시', '출력일시',
                    '브랜드', '배송계약태그', '주문번호', '개별주문번호', '피킹지시서번호', '품고추적번호', 'CS'
                ]
                for tc in req_cols_b2c:
                    if tc not in df_b2c_f.columns: df_b2c_f[tc] = '미지정'
                    df_b2c_f[tc] = df_b2c_f[tc].fillna('미지정')

                if '총출고수량' not in df_b2c_f.columns: df_b2c_f['총출고수량'] = 1
                df_b2c_f['총출고수량'] = pd.to_numeric(df_b2c_f['총출고수량'], errors='coerce').fillna(1)

                valid_mask = ~df_b2c_f['송장번호'].astype(str).str.contains('상세|보기|미지정', na=False)
                df_b2c_valid = df_b2c_f[valid_mask]

                if df_b2c_valid.empty:
                    service.files().update(
                        fileId=file_id, addParents=duplicate_folder_id, removeParents=src_folder,
                        supportsAllDrives=True, fields='id, parents'
                    ).execute()
                    continue

                # 월별 DB 파티셔닝 저장 (wms_b2c_YYYYMM.db)
                df_b2c_valid['month_key'] = df_b2c_valid['영업마감일자'].str.slice(0, 7).str.replace('-', '')
                grouped_months = df_b2c_valid.groupby('month_key')

                for month_key, group_df in grouped_months:
                    month_db_name = f"wms_b2c_{month_key}.db"
                    
                    if not os.path.exists(month_db_name):
                        download_db_from_drive(service, month_db_name)
                        
                    conn_m = sqlite3.connect(month_db_name, timeout=30)

                    conn_m.execute("""
                    CREATE TABLE IF NOT EXISTS shipment_raw (
                        영업마감일자 TEXT, 센터 TEXT, 고객사 TEXT, 배송속성 TEXT, 판매처 TEXT,
                        출고박스종류 TEXT, 송장번호 TEXT, 마감일시 TEXT, 마감자 TEXT, 주문일시 TEXT,
                        결제일시 TEXT, 등록일시 TEXT, 할당일시 TEXT, 출력일시 TEXT, 브랜드 TEXT,
                        배송계약태그 TEXT, 주문번호 TEXT, 개별주문번호 TEXT, 피킹지시서번호 TEXT,
                        품고추적번호 TEXT, CS TEXT, SKU명 TEXT, 바코드 TEXT, 총출고수량 INTEGER,
                        PRIMARY KEY (영업마감일자, 센터, 고객사, 배송속성, 판매처, 출고박스종류, 송장번호, SKU명, 바코드)
                    )
                    """)
                    
                    cursor = conn_m.cursor()
                    cursor.execute("PRAGMA table_info(shipment_raw)")
                    existing_cols = [col[1] for col in cursor.fetchall()]
                    
                    for add_c, col_type in [('SKU명', 'TEXT'), ('바코드', 'TEXT'), ('총출고수량', 'INTEGER')]:
                        if add_c not in existing_cols:
                            try:
                                conn_m.execute(f"ALTER TABLE shipment_raw ADD COLUMN {add_c} {col_type}")
                            except Exception:
                                pass

                    conn_m.execute("""
                    CREATE TABLE IF NOT EXISTS daily_summary (
                        영업마감일자 TEXT, 센터 TEXT, 고객사 TEXT, 배송속성 TEXT, 판매처 TEXT,
                        출고박스종류 TEXT, SKU명 TEXT, 바코드 TEXT, 출고건수 INTEGER, 총출고수량 INTEGER,
                        PRIMARY KEY (영업마감일자, 센터, 고객사, 배송속성, 판매처, 출고박스종류, SKU명, 바코드)
                    )
                    """)

                    shipment_distinct = group_df[
                        ['영업마감일자', '센터', '고객사', '배송속성', '판매처', '출고박스종류', '송장번호',
                         '마감일시', '마감자', '주문일시', '결제일시', '등록일시', '할당일시', '출력일시',
                         '브랜드', '배송계약태그', '주문번호', '개별주문번호', '피킹지시서번호', '품고추적번호', 'CS',
                         'SKU명', '바코드', '총출고수량']
                    ].drop_duplicates(subset=['영업마감일자', '센터', '고객사', '배송속성', '판매처', '출고박스종류', '송장번호', 'SKU명', '바코드'])

                    for _, row_s in shipment_distinct.iterrows():
                        conn_m.execute("""
                        INSERT OR REPLACE INTO shipment_raw
                        (영업마감일자, 센터, 고객사, 배송속성, 판매처, 출고박스종류, 송장번호,
                         마감일시, 마감자, 주문일시, 결제일시, 등록일시, 할당일시, 출력일시,
                         브랜드, 배송계약태그, 주문번호, 개별주문번호, 피킹지시서번호, 품고추적번호, CS,
                         SKU명, 바코드, 총출고수량)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            row_s['영업마감일자'], row_s['센터'], row_s['고객사'], row_s['배송속성'],
                            row_s['판매처'], row_s['출고박스종류'], row_s['송장번호'], str(row_s['마감일시']),
                            str(row_s['마감자']), str(row_s['주문일시']), str(row_s['결제일시']),
                            str(row_s['등록일시']), str(row_s['할당일시']), str(row_s['출력일시']),
                            str(row_s['브랜드']), str(row_s['배송계약태그']), str(row_s['주문번호']),
                            str(row_s['개별주문번호']), str(row_s['피킹지시서번호']), str(row_s['품고추적번호']), str(row_s['CS']),
                            str(row_s['SKU명']), str(row_s['바코드']), int(row_s['총출고수량'])
                        ))

                    b2c_sum = group_df.groupby(
                        ['영업마감일자', '센터', '고객사', '배송속성', '판매처', '출고박스종류', 'SKU명', '바코드']
                    ).agg(
                        출고건수=('송장번호', 'nunique'),
                        총출고수량=('총출고수량', 'sum')
                    ).reset_index()

                    for _, row_b2c in b2c_sum.iterrows():
                        conn_m.execute("""
                        INSERT OR REPLACE INTO daily_summary
                        (영업마감일자, 센터, 고객사, 배송속성, 판매처, 출고박스종류, SKU명, 바코드, 출고건수, 총출고수량)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            row_b2c['영업마감일자'], row_b2c['센터'], row_b2c['고객사'], row_b2c['배송속성'], row_b2c['판매처'],
                            row_b2c['출고박스종류'], row_b2c['SKU명'], row_b2c['바코드'],
                            int(row_b2c['출고건수']), int(row_b2c['총출고수량'])
                        ))

                    conn_m.commit()
                    conn_m.close()
                    b2c_updated_months.add(month_db_name)

                latest_date_str = df_b2c_valid['영업마감일자'].max()
                target_month_str = latest_date_str[:7] if pd.notnull(latest_date_str) and len(str(latest_date_str)) >= 7 else datetime.now().strftime('%Y-%m')
                target_folder_id = get_or_create_subfolder(service, PROCESSED_FOLDER_ID, target_month_str)

                service.files().update(
                    fileId=file_id, addParents=target_folder_id, removeParents=src_folder,
                    supportsAllDrives=True, fields='id, parents'
                ).execute()

        except Exception as file_e:
            print(f"File process error ({file_name}): {file_e}")
            continue

    matched_count, err_msg = 0, ""
    try:
        matched_count, err_msg = update_inbound_plt_box_from_sheets(service, sheets_service, conn_ib)
    except Exception as sheet_e:
        err_msg = str(sheet_e)

    conn_ib.close()

    for month_db_name in b2c_updated_months:
        upload_db_to_drive(service, month_db_name)

    if matched_count > 0 or inbound_updated:
        upload_db_to_drive(service, DB_INBOUND_PATH)

    # ⭐ [안전 연동] B2C 적재 및 업로드 완료 후 구글시트 M/H 데이터 수집 및 DB 동기화
    if sheets_service:
        try:
            sync_mh_to_db(sheets_service, drive_service=service)
        except Exception as mh_e:
            print(f"MH Sync Error: {mh_e}")

    return matched_count, err_msg
